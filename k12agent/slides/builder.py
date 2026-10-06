"""Template mapping: LessonContent -> PowerPoint deck + slide frames (Section 2.1).

* ``build_pptx`` writes an editable .pptx (python-pptx) with the narration in
  the speaker notes and, optionally, the lip-synced avatar clip embedded on
  every slide (bottom-right, as in Figure 8).
* ``render_slide_images`` rasterises the same layout with Pillow; these frames
  are used by the video compositor.
"""

from __future__ import annotations

import re
from pathlib import Path
from typing import List, Optional, Sequence, Tuple

from PIL import Image, ImageDraw, ImageFont

from ..schemas import ContentSection, LessonContent
from .latex import render_formula
from .themes import THEMES, Theme, find_font, pptx_font_name

_TOKENS = re.compile(r"[　-〿㐀-鿿＀-￯]|[^\s　-〿㐀-鿿＀-￯]+|\s+")


def _split_emphasis(text: str, emphasis: Sequence[str]) -> List[Tuple[str, bool]]:
    words = [w for w in emphasis if w]
    if not words:
        return [(text, False)]
    pattern = re.compile("(" + "|".join(re.escape(w) for w in sorted(words, key=len, reverse=True)) + ")")
    return [(part, bool(pattern.fullmatch(part))) for part in pattern.split(text) if part]


# --------------------------------------------------------------------------- #
# PowerPoint
# --------------------------------------------------------------------------- #
def build_pptx(lesson: LessonContent, out_path: str | Path, language: str = "zh",
               avatar_box=(0.66, 0.50, 0.30, 0.46), face_image: Optional[str] = None,
               avatar_clips: Optional[Sequence[Optional[str]]] = None, formula_dir: Optional[str] = None,
               formula_dpi: int = 200) -> Path:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Emu, Pt

    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    formula_dir = Path(formula_dir or out_path.parent / "formulas")
    theme: Theme = THEMES[lesson.stage]
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)  # 16:9
    W, H = prs.slide_width, prs.slide_height
    font = pptx_font_name(language)
    rgb = lambda c: RGBColor(*c)  # noqa: E731

    for i, sec in enumerate(lesson.sections):
        slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
        bg = slide.background.fill
        bg.solid()
        bg.fore_color.rgb = rgb(theme.background)

        label = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, int(W * 0.03), int(H * 0.04), int(W * 0.16), int(H * 0.07))
        label.fill.solid()
        label.fill.fore_color.rgb = rgb(theme.title_bar)
        label.line.fill.background()
        tf = label.text_frame
        tf.text = theme.label.get(language, theme.label["en"])
        tf.paragraphs[0].runs[0].font.size = Pt(16)
        tf.paragraphs[0].runs[0].font.color.rgb = rgb(theme.title_color)
        tf.paragraphs[0].runs[0].font.name = font

        title = slide.shapes.add_textbox(int(W * 0.22), int(H * 0.04), int(W * 0.74), int(H * 0.12))
        title.text_frame.word_wrap = True
        r = title.text_frame.paragraphs[0].add_run()
        r.text = sec.title
        r.font.size, r.font.bold, r.font.name = Pt(34), True, font
        r.font.color.rgb = rgb(theme.text_color)

        body = slide.shapes.add_textbox(int(W * 0.05), int(H * 0.20), int(W * 0.58), int(H * 0.45))
        body.text_frame.word_wrap = True
        for j, bullet in enumerate(sec.bullets):
            para = body.text_frame.paragraphs[0] if j == 0 else body.text_frame.add_paragraph()
            para.space_after = Pt(10)
            for text, emph in _split_emphasis(bullet, sec.emphasis):
                run = para.add_run()
                run.text = text
                run.font.size, run.font.name = Pt(22), font
                run.font.bold = emph
                run.font.color.rgb = rgb(theme.accent if emph else theme.text_color)

        y = int(H * 0.68)
        for f in sec.formulas[:2]:
            try:
                png = render_formula(f.latex, formula_dir, dpi=formula_dpi,
                                     color="#%02x%02x%02x" % theme.text_color)
            except ValueError:
                continue
            pic = slide.shapes.add_picture(str(png), int(W * 0.06), y, height=int(H * 0.11))
            if pic.width > int(W * 0.56):  # keep inside the text column
                ratio = W * 0.56 / pic.width
                pic.width, pic.height = int(pic.width * ratio), int(pic.height * ratio)
            y += pic.height + int(H * 0.02)

        ax, ay, aw, ah = (int(W * avatar_box[0]), int(H * avatar_box[1]), int(W * avatar_box[2]), int(H * avatar_box[3]))
        clip = avatar_clips[i] if avatar_clips and i < len(avatar_clips) else None
        if clip and Path(clip).exists():
            slide.shapes.add_movie(str(clip), ax, ay, aw, ah, poster_frame_image=face_image, mime_type="video/mp4")
        elif face_image and Path(face_image).exists():
            slide.shapes.add_picture(str(face_image), ax, ay, aw, ah)

        slide.notes_slide.notes_text_frame.text = sec.narration
    prs.save(str(out_path))
    return out_path


# --------------------------------------------------------------------------- #
# Raster frames for the video compositor
# --------------------------------------------------------------------------- #
def _wrap(draw: ImageDraw.ImageDraw, text: str, font, max_w: int) -> List[str]:
    lines, cur = [], ""
    for tok in _TOKENS.findall(text):
        trial = cur + tok
        if draw.textlength(trial, font=font) <= max_w or not cur:
            cur = trial
        else:
            lines.append(cur.rstrip())
            cur = tok.lstrip()
    if cur.strip():
        lines.append(cur.rstrip())
    return lines


def _font(path: Optional[str], size: int):
    if path:
        try:
            return ImageFont.truetype(path, size)
        except OSError:
            pass
    try:
        return ImageFont.load_default(size=size)
    except TypeError:  # Pillow < 10.1
        return ImageFont.load_default()


def render_section_image(sec: ContentSection, theme: Theme, size=(1920, 1080), language: str = "zh",
                         font_path: Optional[str] = None, formula_dir: str | Path = "outputs/formulas",
                         reserve_avatar_box=(0.66, 0.50, 0.30, 0.46), subtitle: Optional[str] = None,
                         subtitle_size: int = 32) -> Image.Image:
    W, H = size
    img = Image.new("RGB", size, theme.background)
    d = ImageDraw.Draw(img)
    f_label, f_title, f_body = _font(font_path, int(H * 0.035)), _font(font_path, int(H * 0.055)), _font(font_path, int(H * 0.036))

    d.rectangle([int(W * .03), int(H * .04), int(W * .19), int(H * .11)], fill=theme.title_bar)
    d.text((int(W * .045), int(H * .055)), theme.label.get(language, theme.label["en"]), font=f_label, fill=theme.title_color)
    d.text((int(W * .22), int(H * .045)), sec.title, font=f_title, fill=theme.text_color)

    text_w = int(W * (reserve_avatar_box[0] - 0.08))
    y = int(H * 0.20)
    line_h = int(H * 0.052)
    emph = set(sec.emphasis)
    for bullet in sec.bullets:
        for line in _wrap(d, "• " + bullet if language != "zh" else bullet, f_body, text_w):
            x = int(W * 0.05)
            for part, is_emph in _split_emphasis(line, list(emph)):
                d.text((x, y), part, font=f_body, fill=theme.accent if is_emph else theme.text_color)
                x += int(d.textlength(part, font=f_body))
            y += line_h
        y += int(line_h * 0.35)

    y = max(y + int(H * 0.02), int(H * 0.62))
    for f in sec.formulas[:2]:
        try:
            png = render_formula(f.latex, formula_dir, color="#%02x%02x%02x" % theme.text_color)
        except ValueError:
            continue
        fimg = Image.open(png).convert("RGBA")
        scale = min(int(H * 0.11) / fimg.height, text_w / fimg.width)
        fimg = fimg.resize((max(1, int(fimg.width * scale)), max(1, int(fimg.height * scale))), Image.LANCZOS)
        if y + fimg.height > H * 0.88:
            break
        img.paste(fimg, (int(W * 0.06), y), fimg)
        y += fimg.height + int(H * 0.02)

    if subtitle:
        f_sub = _font(font_path, subtitle_size)
        lines = _wrap(d, subtitle, f_sub, int(W * 0.9))[-2:]
        box_h = (subtitle_size + 12) * len(lines) + 16
        d.rectangle([0, H - box_h, W, H], fill=(0, 0, 0))
        for k, line in enumerate(lines):
            d.text((int(W * 0.05), H - box_h + 8 + k * (subtitle_size + 12)), line, font=f_sub, fill=(255, 255, 255))
    return img


def render_slide_images(lesson: LessonContent, out_dir: str | Path, size=(1920, 1080), language: str = "zh",
                        font_path: Optional[str] = None, avatar_box=(0.66, 0.50, 0.30, 0.46)) -> List[str]:
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    font_path = find_font(font_path)
    theme = THEMES[lesson.stage]
    paths = []
    for i, sec in enumerate(lesson.sections):
        img = render_section_image(sec, theme, size, language, font_path, out_dir / "formulas", avatar_box)
        p = out_dir / f"slide_{i:02d}.png"
        img.save(p)
        paths.append(str(p))
    return paths
